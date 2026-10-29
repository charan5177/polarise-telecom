import sys
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

def create_presentation(output_path="Polarise_Telecom_Presentation.pptx"):
    prs = Presentation()
    # 16:9 Widescreen dimensions
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    
    # Palette
    C_BG_DARK = RGBColor(15, 23, 42)       # #0f172a
    C_CARD_DARK = RGBColor(30, 41, 59)     # #1e293b
    C_CARD_LIGHT = RGBColor(248, 250, 252) # #f8fafc
    C_PRIMARY_BLUE = RGBColor(37, 99, 235) # #2563eb
    C_TEAL = RGBColor(13, 148, 136)        # #0d9488
    C_EMERALD = RGBColor(16, 185, 129)     # #10b981
    C_TEXT_LIGHT = RGBColor(241, 245, 249) # #f1f5f9
    C_TEXT_MUTED = RGBColor(148, 163, 184) # #94a3b8
    C_TEXT_DARK = RGBColor(15, 23, 42)     # #0f172a
    C_BORDER = RGBColor(203, 213, 225)     # #cbd5e1

    def set_slide_background(slide, color):
        background = slide.background
        fill = background.fill
        fill.solid()
        fill.fore_color.rgb = color

    blank_slide_layout = prs.slide_layouts[6]

    # ==========================================
    # SLIDE 1: TITLE SLIDE (Dark Premium Theme)
    # ==========================================
    slide1 = prs.slides.add_slide(blank_slide_layout)
    set_slide_background(slide1, C_BG_DARK)

    # Accent Top Bar
    bar = slide1.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(13.333), Inches(0.12))
    bar.fill.solid()
    bar.fill.fore_color.rgb = C_PRIMARY_BLUE
    bar.line.fill.background()

    # Badge: System Tag
    badge = slide1.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(1.2), Inches(1.3), Inches(3.2), Inches(0.45))
    badge.fill.solid()
    badge.fill.fore_color.rgb = C_CARD_DARK
    badge.line.color.rgb = C_PRIMARY_BLUE
    badge.line.width = Pt(1)
    tf_b = badge.text_frame
    tf_b.text = "CLOUD TELEPHONY & CRM PLATFORM"
    p_b = tf_b.paragraphs[0]
    p_b.font.size = Pt(11)
    p_b.font.bold = True
    p_b.font.color.rgb = C_TEXT_LIGHT
    p_b.alignment = PP_ALIGN.CENTER

    # Main Title
    txBox1 = slide1.shapes.add_textbox(Inches(1.2), Inches(2.0), Inches(10.5), Inches(2.2))
    tf1 = txBox1.text_frame
    tf1.word_wrap = True
    p1 = tf1.paragraphs[0]
    p1.text = "Polarise Telecom Technology"
    p1.font.size = Pt(44)
    p1.font.bold = True
    p1.font.color.rgb = C_TEXT_LIGHT

    p1_sub = tf1.add_paragraph()
    p1_sub.text = "Intelligent Admissions Telecalling Engine & Live Parent CRM Architecture"
    p1_sub.font.size = Pt(20)
    p1_sub.font.color.rgb = C_TEXT_MUTED
    p1_sub.space_before = Pt(12)

    # Bottom Meta Card
    card_meta = slide1.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(1.2), Inches(4.7), Inches(10.9), Inches(1.8))
    card_meta.fill.solid()
    card_meta.fill.fore_color.rgb = C_CARD_DARK
    card_meta.line.color.rgb = RGBColor(51, 65, 85)
    card_meta.line.width = Pt(1)

    tf_meta = card_meta.text_frame
    tf_meta.word_wrap = True
    p_m1 = tf_meta.paragraphs[0]
    p_m1.text = "Client: Polaris Educational Consultancy   •   Target Deployment: Vercel Production"
    p_m1.font.size = Pt(13)
    p_m1.font.bold = True
    p_m1.font.color.rgb = C_TEXT_LIGHT

    p_m2 = tf_meta.add_paragraph()
    p_m2.text = "Infrastructure: FastAPI Serverless (Python 3.12)  |  Vite + React Frontend  |  Exotel Voice PBX Trunk"
    p_m2.font.size = Pt(12)
    p_m2.font.color.rgb = C_TEXT_MUTED
    p_m2.space_before = Pt(6)

    p_m3 = tf_meta.add_paragraph()
    p_m3.text = "Version: 1.0.0 Production Release  •  October 2026 Executive Briefing"
    p_m3.font.size = Pt(11)
    p_m3.font.color.rgb = C_EMERALD
    p_m3.space_before = Pt(8)


    # ==========================================
    # SLIDE 2: BUSINESS CHALLENGE & THE SOLUTION
    # ==========================================
    slide2 = prs.slides.add_slide(blank_slide_layout)
    set_slide_background(slide2, C_CARD_LIGHT)

    # Slide Header
    tb_h2 = slide2.shapes.add_textbox(Inches(0.8), Inches(0.5), Inches(11.5), Inches(1.0))
    tf_h2 = tb_h2.text_frame
    p_h2 = tf_h2.paragraphs[0]
    p_h2.text = "01 | The Business Challenge & Polarise Solution"
    p_h2.font.size = Pt(24)
    p_h2.font.bold = True
    p_h2.font.color.rgb = C_TEXT_DARK

    p_h2_sub = tf_h2.add_paragraph()
    p_h2_sub.text = "Transforming manual, vulnerable telecalling into a secure, carrier-bridged admissions powerhouse"
    p_h2_sub.font.size = Pt(13)
    p_h2_sub.font.color.rgb = RGBColor(100, 116, 139)

    # 3 Column Cards
    col_w = Inches(3.64)
    gap = Inches(0.35)
    y_card = Inches(1.7)
    card_h = Inches(5.0)

    cards_data = [
        ("The Challenge", RGBColor(239, 68, 68), [
            ("Unprotected Counselor Privacy", "Counselors previously dialed parents from personal mobile numbers, resulting in unwanted callbacks and lost privacy."),
            ("Lead Conflicts & Duplication", "Manual Excel sheets led to overlapping calls, duplicate inquiries, and student parent dissatisfaction."),
            ("No Quality Supervision", "Absence of call audio recordings made it impossible to evaluate counseling pitch compliance or resolve student disputes.")
        ]),
        ("The Polarise Solution", C_PRIMARY_BLUE, [
            ("Virtual Caller ID Masking", "Every outbound call displays the official Polaris virtual number (+918047190000), protecting staff confidentiality."),
            ("Dedicated 100-200 Queues", "Admin divides bulk student leads across 10 counselors with automated, collision-free queue allocation."),
            ("Mandatory Call Dispositions", "Structured call-wrap notes and outcome tracking (Interested, Callback, Admitted) ensure full pipeline visibility.")
        ]),
        ("Core Capabilities", C_EMERALD, [
            ("10-Digit Phone Number Guard", "Strict front-end validation prevents dialing invalid or truncated student contact numbers."),
            ("Admin Audio Recording Vault", "Audio records every call automatically. Recordings are strictly restricted to administrators for quality compliance."),
            ("Real-Time Agent Monitoring", "Supervisors track live staff statuses (Available, On Call, Break, Wrap, Offline) in real time.")
        ])
    ]

    for i, (title, accent, bullets) in enumerate(cards_data):
        x = Inches(0.8) + i * (col_w + gap)
        card = slide2.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, y_card, col_w, card_h)
        card.fill.solid()
        card.fill.fore_color.rgb = RGBColor(255, 255, 255)
        card.line.color.rgb = C_BORDER
        card.line.width = Pt(1)

        # Top Accent Line on Card
        top_line = slide2.shapes.add_shape(MSO_SHAPE.RECTANGLE, x, y_card, col_w, Inches(0.1))
        top_line.fill.solid()
        top_line.fill.fore_color.rgb = accent
        top_line.line.fill.background()

        tf = card.text_frame
        tf.word_wrap = True
        tf.margin_left = Inches(0.25)
        tf.margin_right = Inches(0.25)
        tf.margin_top = Inches(0.3)

        p = tf.paragraphs[0]
        p.text = title
        p.font.size = Pt(16)
        p.font.bold = True
        p.font.color.rgb = accent

        for b_title, b_desc in bullets:
            pb_t = tf.add_paragraph()
            pb_t.text = f"•  {b_title}"
            pb_t.font.size = Pt(12)
            pb_t.font.bold = True
            pb_t.font.color.rgb = C_TEXT_DARK
            pb_t.space_before = Pt(12)

            pb_d = tf.add_paragraph()
            pb_d.text = b_desc
            pb_d.font.size = Pt(10.5)
            pb_d.font.color.rgb = RGBColor(71, 85, 105)
            pb_d.space_before = Pt(3)


    # ==========================================
    # SLIDE 3: TELEPHONY ARCHITECTURE & CONFIG
    # ==========================================
    slide3 = prs.slides.add_slide(blank_slide_layout)
    set_slide_background(slide3, C_CARD_LIGHT)

    tb_h3 = slide3.shapes.add_textbox(Inches(0.8), Inches(0.5), Inches(11.5), Inches(1.0))
    tf_h3 = tb_h3.text_frame
    p_h3 = tf_h3.paragraphs[0]
    p_h3.text = "02 | Cloud Telephony Architecture & TelephonyConfig"
    p_h3.font.size = Pt(24)
    p_h3.font.bold = True
    p_h3.font.color.rgb = C_TEXT_DARK

    p_h3_sub = tf_h3.add_paragraph()
    p_h3_sub.text = "How the TelephonyConfig adapter bridges CRM clicks to real telecom voice lines"
    p_h3_sub.font.size = Pt(13)
    p_h3_sub.font.color.rgb = RGBColor(100, 116, 139)

    # 4 Architecture Blocks (2x2 Grid)
    grid_w = Inches(5.6)
    grid_h = Inches(2.4)
    grid_x1 = Inches(0.8)
    grid_x2 = Inches(6.8)
    grid_y1 = Inches(1.7)
    grid_y2 = Inches(4.4)

    blocks = [
        (grid_x1, grid_y1, "1. Click-to-Call PBX Bridging", C_PRIMARY_BLUE,
         "When a counselor clicks 'Initiate Call':\n"
         "1. Polarise backend invokes Exotel REST API.\n"
         "2. PBX first rings the counselor's phone / extension.\n"
         "3. As soon as the counselor answers, PBX dials the parent.\n"
         "4. Lines bridge seamlessly with 0 setup time."),

        (grid_x2, grid_y1, "2. Virtual Number Masking (Caller ID)", C_TEAL,
         "Institutional Privacy Protection:\n"
         "• Outbound Caller ID: +918047190000 (Polaris Virtual Number).\n"
         "• Counselor personal numbers are completely shielded.\n"
         "• Parents recognize official consultancy branding.\n"
         "• Eliminates personal off-hours phone disturbance."),

        (grid_x1, grid_y2, "3. Dual-Mode Telephony Engine", C_EMERALD,
         "Seamless Toggle between Environments:\n"
         "• Demo Sandbox Mode: Generates synthetic WAV waveforms & simulated call states for risk-free staff training at $0.\n"
         "• Live Exotel Mode: Full carrier integration via Account SID, API Key, Token, and Subdomain with 1-click admin toggle."),

        (grid_x2, grid_y2, "4. Webhooks & Call Forwarding", RGBColor(147, 51, 234),
         "Automated Carrier Synchronization:\n"
         "• Webhook receiver (/api/telephony/webhooks/exotel) updates call duration, recording URLs, and status without polling.\n"
         "• Call Forwarding: Inbound parent calls route to available counselors or designated priority desks automatically.")
    ]

    for x, y, title, accent, content in blocks:
        block = slide3.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, y, grid_w, grid_h)
        block.fill.solid()
        block.fill.fore_color.rgb = RGBColor(255, 255, 255)
        block.line.color.rgb = C_BORDER
        block.line.width = Pt(1)

        # Accent Bar
        top_b = slide3.shapes.add_shape(MSO_SHAPE.RECTANGLE, x, y, grid_w, Inches(0.08))
        top_b.fill.solid()
        top_b.fill.fore_color.rgb = accent
        top_b.line.fill.background()

        tf = block.text_frame
        tf.word_wrap = True
        tf.margin_left = Inches(0.25)
        tf.margin_right = Inches(0.25)
        tf.margin_top = Inches(0.2)

        p = tf.paragraphs[0]
        p.text = title
        p.font.size = Pt(14)
        p.font.bold = True
        p.font.color.rgb = accent

        p_c = tf.add_paragraph()
        p_c.text = content
        p_c.font.size = Pt(10.5)
        p_c.font.color.rgb = RGBColor(51, 65, 85)
        p_c.space_before = Pt(6)


    # ==========================================
    # SLIDE 4: ROLES & OPERATIONAL WORKFLOWS
    # ==========================================
    slide4 = prs.slides.add_slide(blank_slide_layout)
    set_slide_background(slide4, C_CARD_LIGHT)

    tb_h4 = slide4.shapes.add_textbox(Inches(0.8), Inches(0.5), Inches(11.5), Inches(1.0))
    tf_h4 = tb_h4.text_frame
    p_h4 = tf_h4.paragraphs[0]
    p_h4.text = "03 | User Roles & Daily Operational Workflows"
    p_h4.font.size = Pt(24)
    p_h4.font.bold = True
    p_h4.font.color.rgb = C_TEXT_DARK

    p_h4_sub = tf_h4.add_paragraph()
    p_h4_sub.text = "Role-based separation of concerns between Admissions Counselors and Administrative Operations"
    p_h4_sub.font.size = Pt(13)
    p_h4_sub.font.color.rgb = RGBColor(100, 116, 139)

    # 2 Large Columns: Counselor vs Admin
    half_w = Inches(5.6)
    half_h = Inches(5.1)
    y_half = Inches(1.7)

    # Left: Admissions Counselor Station
    c_left = slide4.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), y_half, half_w, half_h)
    c_left.fill.solid()
    c_left.fill.fore_color.rgb = RGBColor(255, 255, 255)
    c_left.line.color.rgb = C_BORDER
    c_left.line.width = Pt(1)

    t_left = slide4.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.8), y_half, half_w, Inches(0.1))
    t_left.fill.solid()
    t_left.fill.fore_color.rgb = C_PRIMARY_BLUE
    t_left.line.fill.background()

    tf_l = c_left.text_frame
    tf_l.word_wrap = True
    tf_l.margin_left = Inches(0.3)
    tf_l.margin_right = Inches(0.3)
    tf_l.margin_top = Inches(0.25)

    pl_t = tf_l.paragraphs[0]
    pl_t.text = "Admissions Counselor Portal (10 Staff)"
    pl_t.font.size = Pt(15)
    pl_t.font.bold = True
    pl_t.font.color.rgb = C_PRIMARY_BLUE

    items_counselor = [
        ("Dedicated Lead Pipeline", "Counselors receive 100 to 200 student leads assigned by admin. They only see their own queue with zero cross-leakage."),
        ("10-Digit Phone Number Guard", "Strict validation blocks accidental dialing if input numbers exceed or lack 10 digits."),
        ("Click-to-Dial & Active Call Timer", "Instant carrier connection with live duration timer and parent profile details displayed onscreen."),
        ("Mandatory Call-Wrap Notes", "Requires counselors to record student interest and choose call disposition before moving to the next lead."),
        ("Strict Audio Privacy", "Call recordings are completely hidden from employees to ensure compliance and institutional governance.")
    ]
    for it, idesc in items_counselor:
        p1 = tf_l.add_paragraph()
        p1.text = f"✔  {it}"
        p1.font.size = Pt(11.5)
        p1.font.bold = True
        p1.font.color.rgb = C_TEXT_DARK
        p1.space_before = Pt(8)

        p2 = tf_l.add_paragraph()
        p2.text = idesc
        p2.font.size = Pt(10)
        p2.font.color.rgb = RGBColor(71, 85, 105)
        p2.space_before = Pt(2)

    # Right: Administrative Command
    c_right = slide4.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(6.8), y_half, half_w, half_h)
    c_right.fill.solid()
    c_right.fill.fore_color.rgb = RGBColor(255, 255, 255)
    c_right.line.color.rgb = C_BORDER
    c_right.line.width = Pt(1)

    t_right = slide4.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(6.8), y_half, half_w, Inches(0.1))
    t_right.fill.solid()
    t_right.fill.fore_color.rgb = C_TEAL
    t_right.line.fill.background()

    tf_r = c_right.text_frame
    tf_r.word_wrap = True
    tf_r.margin_left = Inches(0.3)
    tf_r.margin_right = Inches(0.3)
    tf_r.margin_top = Inches(0.25)

    pr_t = tf_r.paragraphs[0]
    pr_t.text = "Operations & Admin Command (3 Roles)"
    pr_t.font.size = Pt(15)
    pr_t.font.bold = True
    pr_t.font.color.rgb = C_TEAL

    items_admin = [
        ("Live Agent Supervision Board", "Real-time indicators showing available, on-call, wrap-up, and offline counselors with live extension monitoring."),
        ("Bulk Excel Lead Ingestion", "Uploads hundreds of leads via .xlsx/.csv, automatically strips cost/pricing, and distributes 100-200 leads to selected counselors."),
        ("Audio Recordings Vault", "Secure audio station featuring HTML5 audio player with waveform scrubber to review counselor pitch quality."),
        ("Day / Week / Month Analytics", "Dynamic telemetry computed from real call logs with multi-colored bar charts and chromatic gradient spline curves."),
        ("Locked Profile Security", "Admins can lock employee accounts; locked profiles display a bold 'Contact Administrator' security blocker on login.")
    ]
    for it, idesc in items_admin:
        p1 = tf_r.add_paragraph()
        p1.text = f"✔  {it}"
        p1.font.size = Pt(11.5)
        p1.font.bold = True
        p1.font.color.rgb = C_TEXT_DARK
        p1.space_before = Pt(8)

        p2 = tf_r.add_paragraph()
        p2.text = idesc
        p2.font.size = Pt(10)
        p2.font.color.rgb = RGBColor(71, 85, 105)
        p2.space_before = Pt(2)


    # ==========================================
    # SLIDE 5: PRODUCTION & MEASURABLE IMPACT
    # ==========================================
    slide5 = prs.slides.add_slide(blank_slide_layout)
    set_slide_background(slide5, C_BG_DARK)

    # Slide Header
    tb_h5 = slide5.shapes.add_textbox(Inches(0.8), Inches(0.5), Inches(11.5), Inches(1.0))
    tf_h5 = tb_h5.text_frame
    p_h5 = tf_h5.paragraphs[0]
    p_h5.text = "04 | Production Deployment & Business Impact"
    p_h5.font.size = Pt(24)
    p_h5.font.bold = True
    p_h5.font.color.rgb = C_TEXT_LIGHT

    p_h5_sub = tf_h5.add_paragraph()
    p_h5_sub.text = "Verified live deployment on Vercel with measurable operational benefits"
    p_h5_sub.font.size = Pt(13)
    p_h5_sub.font.color.rgb = C_TEXT_MUTED

    # Top Metric Banners
    met_w = Inches(2.7)
    met_h = Inches(1.3)
    met_y = Inches(1.6)

    metrics = [
        ("100%", "Phone Privacy", "Zero counselor personal numbers exposed", C_PRIMARY_BLUE),
        ("0 Leads", "Collision / Conflict", "Strict isolated counselor allocations", C_TEAL),
        ("3x", "Calling Throughput", "Click-to-call PBX eliminates dial lag", C_EMERALD),
        ("100%", "Call Auditing", "Secure admin audio vault archiving", RGBColor(168, 85, 247))
    ]

    for i, (m_val, m_lbl, m_sub, m_clr) in enumerate(metrics):
        x = Inches(0.8) + i * (met_w + Inches(0.28))
        m_card = slide5.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, met_y, met_w, met_h)
        m_card.fill.solid()
        m_card.fill.fore_color.rgb = C_CARD_DARK
        m_card.line.color.rgb = m_clr
        m_card.line.width = Pt(1)

        tf = m_card.text_frame
        tf.word_wrap = True
        tf.margin_top = Inches(0.12)
        p1 = tf.paragraphs[0]
        p1.text = m_val
        p1.font.size = Pt(26)
        p1.font.bold = True
        p1.font.color.rgb = m_clr
        p1.alignment = PP_ALIGN.CENTER

        p2 = tf.add_paragraph()
        p2.text = f"{m_lbl} — {m_sub}"
        p2.font.size = Pt(9.5)
        p2.font.color.rgb = C_TEXT_MUTED
        p2.alignment = PP_ALIGN.CENTER

    # Bottom Deployment Specs Card
    dep_card = slide5.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(3.2), Inches(11.733), Inches(3.6))
    dep_card.fill.solid()
    dep_card.fill.fore_color.rgb = C_CARD_DARK
    dep_card.line.color.rgb = RGBColor(51, 65, 85)
    dep_card.line.width = Pt(1)

    tf_d = dep_card.text_frame
    tf_d.word_wrap = True
    tf_d.margin_left = Inches(0.4)
    tf_d.margin_right = Inches(0.4)
    tf_d.margin_top = Inches(0.3)

    pd_t = tf_d.paragraphs[0]
    pd_t.text = "Live Production Infrastructure & Master Access"
    pd_t.font.size = Pt(16)
    pd_t.font.bold = True
    pd_t.font.color.rgb = C_TEXT_LIGHT

    specs = [
        ("Live Web Application URL:", "https://polarise-telecom.vercel.app  (Vercel Edge Global CDN • SSL Secured)"),
        ("Synchronized GitHub Repository:", "https://github.com/charan5177/polarise-telecom  (Continuous Deployment Pipeline)"),
        ("Backend Serverless Health:", "https://polarise-telecom.vercel.app/api/health  (FastAPI • 200 OK Active)"),
        ("Super Administrator Login:", "admin1@polarisedu.in  |  Password: Polaris@2026"),
        ("Operations Administrator Login:", "admin2@polarisedu.in  |  Password: Polaris@2026"),
        ("Admissions Counselor Logins:", "emp01@polarisedu.in through emp10@polarisedu.in  |  Password: Polaris@2026")
    ]

    for label, val in specs:
        p_row = tf_d.add_paragraph()
        p_row.text = f"•  {label}  {val}"
        p_row.font.size = Pt(11)
        p_row.font.color.rgb = C_TEXT_MUTED
        p_row.space_before = Pt(7)

    prs.save(output_path)
    print(f"PowerPoint Presentation created successfully at: {output_path}")

if __name__ == "__main__":
    out = sys.argv[1] if len(sys.argv) > 1 else "Polarise_Telecom_Presentation.pptx"
    create_presentation(out)
